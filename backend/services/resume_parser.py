import os
import re
from pathlib import Path

class ResumeParser:
    @staticmethod
    def extract_text(file_path: str) -> str:
        path = Path(file_path)
        if not path.exists():
            raise FileNotFoundError(f"File not found: {file_path}")
            
        ext = path.suffix.lower()
        if ext == ".pdf":
            return ResumeParser._extract_pdf(path)
        elif ext in (".docx", ".doc"):
            return ResumeParser._extract_docx(path)
        elif ext == ".txt":
            with open(path, "r", encoding="utf-8", errors="ignore") as f:
                return f.read()
        else:
            raise ValueError(f"Unsupported file extension: {ext}. Allowed: PDF, DOCX, TXT")

    @staticmethod
    def _extract_pdf(path: Path) -> str:
        text_parts = []
        # Try pypdf first
        try:
            import pypdf
            reader = pypdf.PdfReader(str(path))
            for page in reader.pages:
                extracted = page.extract_text()
                if extracted:
                    text_parts.append(extracted)
            if text_parts:
                return "\n".join(text_parts).strip()
        except Exception:
            pass

        # Try PyPDF2 fallback
        try:
            import PyPDF2
            with open(path, "rb") as f:
                r = PyPDF2.PdfReader(f)
                for page in r.pages:
                    t = page.extract_text()
                    if t:
                        text_parts.append(t)
            if text_parts:
                return "\n".join(text_parts).strip()
        except Exception:
            pass

        # Try pdfplumber fallback
        try:
            import pdfplumber
            with pdfplumber.open(str(path)) as pdf:
                for p in pdf.pages:
                    t = p.extract_text()
                    if t:
                        text_parts.append(t)
            if text_parts:
                return "\n".join(text_parts).strip()
        except Exception:
            pass

        if not text_parts:
            raise ValueError("Unable to parse text from PDF. The document may be scanned or empty.")
        return "\n".join(text_parts).strip()

    @staticmethod
    def _extract_docx(path: Path) -> str:
        try:
            import docx
            doc = docx.Document(str(path))
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            paragraphs.append(cell.text.strip())
            return "\n".join(paragraphs).strip()
        except Exception as e:
            raise ValueError(f"Unable to parse DOCX document: {e}")

    @staticmethod
    def parse(file_path: str = None, raw_text: str = None) -> dict:
        text = raw_text or ""
        if file_path:
            text = ResumeParser.extract_text(file_path)
            
        if not text.strip():
            raise ValueError("Resume is empty or contains unreadable content.")
            
        lines = [line.strip() for line in text.split("\n") if line.strip()]
        
        candidate_name = "Candidate"
        if lines:
            first_line = lines[0]
            if len(first_line.split()) <= 4 and not re.search(r"[@\d]", first_line):
                candidate_name = first_line.title()
                
        email_match = re.search(r"[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+", text)
        phone_match = re.search(r"(\+?\d{1,3}[-.\s]?)?\(?\d{3}\)?[-.\s]?\d{3}[-.\s]?\d{4}", text)
        
        sections = ResumeParser._segment_sections(text)
        
        return {
            "candidate_name": candidate_name,
            "email": email_match.group(0) if email_match else None,
            "phone": phone_match.group(0) if phone_match else None,
            "full_text": text,
            "sections": sections
        }

    @staticmethod
    def _segment_sections(text: str) -> dict:
        section_headers = {
            "skills": r"(technical\s+skills|skills|technologies|proficiencies|core\s+competencies)",
            "experience": r"(work\s+experience|experience|employment|internships|internship\s+experience)",
            "projects": r"(projects|academic\s+projects|personal\s+projects)",
            "education": r"(education|academic\s+background|qualifications)",
            "certifications": r"(certifications|certificates|achievements)"
        }
        
        sections = {k: "" for k in section_headers}
        current_section = "skills"
        
        for line in text.split("\n"):
            stripped = line.strip()
            matched_header = None
            for key, pattern in section_headers.items():
                if re.match(rf"^{pattern}[:\s]*$", stripped, re.IGNORECASE):
                    matched_header = key
                    break
            if matched_header:
                current_section = matched_header
            else:
                sections[current_section] += line + "\n"
                
        return sections
